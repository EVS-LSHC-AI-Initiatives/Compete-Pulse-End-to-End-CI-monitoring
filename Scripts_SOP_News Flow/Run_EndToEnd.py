"""
Run_EndToEnd.py
----------------
Chains the four MASH CI monitoring stages together in one run:

  1. RSS extraction    -- Run_All_RSS.py (interactive: pick pipelines + date range)
                           -> Output_Combined/Combined_Report_<timestamp>.xlsx
  2. Link deduplication -- dedupe_source_urls.py (resolves Google redirect
                           URLs, drops duplicate articles)
                           -> Deduped_Links/<...>_Deduped.xlsx
  3. Parallel AI triage -- parallel_triage_agent.py (Task Group API)
                           -> Output_Triage_Agent/<...>_Triaged_<timestamp>.xlsx
  4. Mail drafting      -- draft_morning_update.py (Task Group API; pre-filter,
                           provenance check, then drafting -- see its own
                           module docstring for the selection pipeline)
                           -> Output_Mail_Draft/<slot>_Update_<date>_<timestamp>.docx

Every script here is a self-contained copy living next to this one -- this
folder does not depend on anything outside itself except your Parallel API
key, which you must paste into "Parallel API Key.txt" in this same folder
before running stage 3 or 4.

Stage 4 spends additional Parallel AI credit on top of stage 3 (a provenance
check plus a drafting call per surviving candidate) on every run. It finding
zero mail-worthy items is a normal outcome, not a pipeline failure -- Stage 4
failing does not abort the run or mark stages 1-3's output as bad; it just
means no docx was produced this time. Pass --skip-mail-draft to stop after
triage (stage 3) and skip stage 4 entirely.

Usage:
    python Run_EndToEnd.py
        (double-click, or run with no arguments -- stage 1 will prompt you
        for which RSS pipelines to run and a date range, same as always;
        stages 2-4 then run automatically with no further input)
    python Run_EndToEnd.py --processor core
        (use a different Parallel processor tier for triage AND drafting)
    python Run_EndToEnd.py --skip-mail-draft
        (stop after triage; do not run stage 4)
"""

import argparse
import builtins
import codecs
import json
import os
import subprocess
import sys
import time

import openpyxl

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(SCRIPT_DIR)

RUN_ALL_RSS = os.path.join(SCRIPT_DIR, "Run_All_RSS.py")
DEDUPE_SCRIPT = os.path.join(SCRIPT_DIR, "dedupe_source_urls.py")
TRIAGE_SCRIPT = os.path.join(SCRIPT_DIR, "parallel_triage_agent.py")
MAIL_SCRIPT = os.path.join(SCRIPT_DIR, "draft_morning_update.py")
API_KEY_FILE = os.path.join(SCRIPT_DIR, "Parallel API Key.txt")
# Written by Run_All_RSS.py with the START of the date filter you entered at
# its prompt, so stage 4 can use that as the exact provenance cutoff instead
# of guessing one from the slot's lookback hours. See read_extraction_window_start().
WINDOW_STATE_FILE = os.path.join(SCRIPT_DIR, "_extraction_window_state.json")

OUTPUT_COMBINED_DIR = os.path.join(ROOT_DIR, "Output_Combined")
DEDUPED_LINKS_DIR = os.path.join(ROOT_DIR, "Deduped_Links")
OUTPUT_TRIAGE_DIR = os.path.join(ROOT_DIR, "Output_Triage_Agent")
OUTPUT_MAIL_DRAFT_DIR = os.path.join(ROOT_DIR, "Output_Mail_Draft")
LOGS_DIR = os.path.join(ROOT_DIR, "Logs")

# rss_opml_common.py's run_pipeline() writes straight to these folders
# without ever creating them -- it only "works" in the original RSS_2 July
# folder because they already exist there from an earlier run. This fresh
# workflow folder needs them created up front.
EXTRACTION_OUTPUT_DIRS = [
    "Output_ClinicalTrials",
    "Output_GoogleAlerts",
    "Output_CompanyAlerts",
    "Out_MediaAggregators",
    "Out_RegulatoryAlerts",
    "Output_Combined",
    "Logs",
]


def ensure_extraction_output_dirs():
    for name in EXTRACTION_OUTPUT_DIRS:
        os.makedirs(os.path.join(ROOT_DIR, name), exist_ok=True)


def check_api_key_preflight():
    """Warn early if the Parallel API key looks unset, before spending
    10+ minutes on RSS extraction only to fail at the final stage."""
    key = ""
    if os.path.exists(API_KEY_FILE):
        with open(API_KEY_FILE, "r", encoding="utf-8") as f:
            key = f.read().strip()
    if key:
        return
    print("=" * 70)
    print("  WARNING: 'Parallel API Key.txt' is empty or missing.")
    print(f"  Path: {API_KEY_FILE}")
    print("  Stage 3 (triage) will fail without it.")
    print("=" * 70)
    if sys.stdin.isatty():
        choice = input("Continue anyway with stages 1-2 only? (y/N): ").strip().lower()
        if choice != "y":
            sys.exit("Aborted. Paste your Parallel API key into the file above and re-run.")


def read_extraction_window_start():
    """Read the START date/time entered at Run_All_RSS.py's date-filter prompt
    during this run's stage 1, as written to WINDOW_STATE_FILE. Returns the
    'YYYY-MM-DD HH:MM' string to pass as draft_morning_update.py's
    --window-start, or None if stage 1's prompt was skipped (no filter
    entered) or the state file is missing/unreadable -- in which case stage 4
    falls back to its own lookback-hours default, same as before."""
    if not os.path.exists(WINDOW_STATE_FILE):
        return None
    try:
        with open(WINDOW_STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f).get("start_dt")
    except (OSError, ValueError):
        return None


def newest_xlsx(directory, prefix=None, exclude_substrings=()):
    if not os.path.isdir(directory):
        return None
    candidates = []
    for name in os.listdir(directory):
        if not name.endswith(".xlsx"):
            continue
        if prefix and not name.startswith(prefix):
            continue
        if any(substr in name for substr in exclude_substrings):
            continue
        candidates.append(os.path.join(directory, name))
    if not candidates:
        return None
    candidates.sort(key=os.path.getmtime, reverse=True)
    return candidates[0]


def read_triage_insights_counts(xlsx_path):
    """Read the Priority-breakdown counts from a triage output's Insights
    sheet, so we can tell a run that's genuinely complete from one that gave
    up on some rows (ERROR / NOT PROCESSED) without raising an exception."""
    if not xlsx_path or not os.path.exists(xlsx_path):
        return {}
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if "Insights" not in wb.sheetnames:
        return {}
    ws = wb["Insights"]
    counts = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] in ("High", "Medium", "Low", "ERROR", "NOT PROCESSED"):
            counts[row[0]] = row[1] or 0
    return counts


class TeeLogger:
    """Tee all stream writes to both the console stream and an active log file."""

    def __init__(self, original_stream, log_file):
        self.original_stream = original_stream
        self.log_file = log_file

    def write(self, data):
        try:
            self.original_stream.write(data)
            self.original_stream.flush()
        except Exception:
            pass
        if self.log_file and not self.log_file.closed:
            try:
                self.log_file.write(data)
                self.log_file.flush()
            except Exception:
                pass

    def flush(self):
        try:
            self.original_stream.flush()
        except Exception:
            pass
        if self.log_file and not self.log_file.closed:
            try:
                self.log_file.flush()
            except Exception:
                pass

    def isatty(self):
        return hasattr(self.original_stream, "isatty") and self.original_stream.isatty()

    def fileno(self):
        return self.original_stream.fileno()


_original_input = builtins.input


def make_logged_input(log_file_container):
    def logged_input(prompt=""):
        val = _original_input(prompt)
        f = log_file_container[0] if log_file_container else None
        if f and not f.closed:
            try:
                f.write(f"{val}\n")
                f.flush()
            except Exception:
                pass
        return val

    return logged_input


def run_stage(label, cmd, fatal=True):
    """Run one stage as a subprocess. By default a non-zero exit aborts the
    whole end-to-end run (fatal=True) -- appropriate for stages 1-3, whose
    output every later stage depends on. Stage 4 passes fatal=False: it
    legitimately exits non-zero when zero items are mail-worthy this run,
    which is a normal outcome, not a broken pipeline, and must not throw away
    the already-valid Stage 1-3 output. Output is streamed live to console
    and mirrored into the run log."""
    print("\n" + "=" * 70)
    print(f"  STAGE: {label}")
    print("=" * 70)
    t0 = time.time()
    # Force UTF-8 stdout/stderr in the child so a non-ASCII character in a
    # feed label, article title, or summary (e.g. a Greek letter in a drug
    # name) can't crash the stage with a UnicodeEncodeError under the
    # Windows console's default cp1252 codepage. ENDTOEND_ORCHESTRATED tells
    # each stage to skip its own "Press Enter to close" pause -- otherwise
    # the run needs a keypress after every stage instead of running straight
    # through to the end. PYTHONUNBUFFERED ensures prompts and logs stream
    # immediately through the pipe.
    child_env = dict(
        os.environ,
        PYTHONIOENCODING="utf-8",
        PYTHONUNBUFFERED="1",
        ENDTOEND_ORCHESTRATED="1",
    )
    p = subprocess.Popen(
        cmd,
        cwd=SCRIPT_DIR,
        env=child_env,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    decoder = codecs.getincrementaldecoder("utf-8")(errors="replace")
    try:
        while True:
            chunk = os.read(p.stdout.fileno(), 4096)
            if not chunk:
                break
            text = decoder.decode(chunk)
            sys.stdout.write(text)
            sys.stdout.flush()
        rest = decoder.decode(b"", final=True)
        if rest:
            sys.stdout.write(rest)
            sys.stdout.flush()
        p.wait()
    except KeyboardInterrupt:
        p.kill()
        p.wait()
        raise

    elapsed = time.time() - t0
    if p.returncode != 0:
        message = f"\n[{label}] exited with code {p.returncode}."
        if fatal:
            sys.exit(message + " -- stopping end-to-end run.")
        print(message + " -- continuing (this stage's output is optional).")
        return elapsed, False
    return elapsed, True


def main():
    parser = argparse.ArgumentParser(description="Run the full RSS extraction -> dedup -> triage -> mail-draft pipeline")
    parser.add_argument("--processor", choices=["lite", "base", "core"], default="base", help="Parallel processor tier for the triage AND mail-drafting stages")
    parser.add_argument("--skip-mail-draft", action="store_true", help="Stop after triage (stage 3); do not run the mail-drafting stage")
    args = parser.parse_args()

    ensure_extraction_output_dirs()
    timestamp = time.strftime("%Y-%m-%d_%H%M%S")
    log_filename = f"Run_EndToEnd_{timestamp}.txt"
    log_path = os.path.join(LOGS_DIR, log_filename)
    log_file = open(log_path, "w", encoding="utf-8", errors="replace", newline="")

    orig_stdout = sys.stdout
    orig_stderr = sys.stderr
    log_container = [log_file]
    sys.stdout = TeeLogger(orig_stdout, log_file)
    sys.stderr = TeeLogger(orig_stderr, log_file)
    builtins.input = make_logged_input(log_container)

    exit_code = 0
    try:
        print(f"[Logging] Run started. Log file: {log_path}")

        # ── Ask up front, before Stage 1's own prompt, so the whole run (including
        # Stage 4's output format) is decided before you walk away from the
        # keyboard. Skipped entirely for --skip-mail-draft (stage 4 never runs)
        # and for unattended/non-tty invocations (defaults: proceed, docx).
        proceed_with_draft = True
        draft_output_format = "docx"
        if not args.skip_mail_draft and sys.stdin.isatty():
            answer = input(
                "\nProceed with the provenance check and drafting step (Stage 4) once triage finishes? "
                "(y/n, default y): "
            ).strip().lower()
            if answer == "n":
                proceed_with_draft = False
                print("Stage 4 will run in dry-run mode: pre-filter only, no provenance check, "
                      "no drafting, no additional API cost.")
            else:
                print("\nDrafting output:")
                word_answer = input("Create a Word document (.docx)? (y/n, default y): ").strip().lower()
                want_word = word_answer != "n"
                outlook_answer = input("Create an Outlook draft? (y/n, default y): ").strip().lower()
                want_outlook_choice = outlook_answer != "n"
                if want_word and want_outlook_choice:
                    draft_output_format = "both"
                elif want_outlook_choice:
                    draft_output_format = "outlook"
                elif want_word:
                    draft_output_format = "docx"
                else:
                    draft_output_format = "docx"
                    print("Neither was selected; defaulting to Word document so a draft is still produced.")
                print(f"Output format: {draft_output_format.upper()}")

        check_api_key_preflight()

        overall_start = time.time()
        total_stages = 3 if args.skip_mail_draft else 4

        # Stage 1: RSS extraction (interactive menu + date range). Clear any stale
        # window-state file up front so a crashed/skipped stage 1 can't leave a
        # PREVIOUS run's date filter lying around for stage 4 to pick up by mistake.
        try:
            os.remove(WINDOW_STATE_FILE)
        except FileNotFoundError:
            pass
        before_combined = newest_xlsx(OUTPUT_COMBINED_DIR, prefix="Combined_Report_")
        before_mtime = os.path.getmtime(before_combined) if before_combined else 0
        stage1_elapsed, _ = run_stage(f"1/{total_stages} RSS Extraction (Run_All_RSS.py)", [sys.executable, RUN_ALL_RSS])

        combined_report = newest_xlsx(OUTPUT_COMBINED_DIR, prefix="Combined_Report_")
        if not combined_report or os.path.getmtime(combined_report) <= before_mtime:
            sys.exit(
                "\nNo new Combined_Report produced by stage 1 (no pipelines selected, or no data found). "
                "Stopping -- nothing to deduplicate or triage."
            )
        print(f"\n[Stage 1] Combined report: {combined_report}")

        # The START of the date filter you entered at stage 1's prompt, if any --
        # forwarded to stage 4 below as --window-start so the provenance cutoff
        # matches the window you actually asked for instead of a lookback guess.
        window_start = read_extraction_window_start()
        if window_start:
            print(f"[Stage 1] Extraction window start: {window_start} (will be used as stage 4's provenance cutoff)")

        # Stage 2: link dedup
        stage2_elapsed, _ = run_stage(
            f"2/{total_stages} Link Deduplication (dedupe_source_urls.py)",
            [sys.executable, DEDUPE_SCRIPT, "--input", combined_report],
        )
        base, ext = os.path.splitext(os.path.basename(combined_report))
        deduped_report = os.path.join(DEDUPED_LINKS_DIR, f"{base}_Deduped{ext}")
        if not os.path.exists(deduped_report):
            sys.exit(f"\n[Stage 2] Expected output not found: {deduped_report}")
        print(f"\n[Stage 2] Deduped report: {deduped_report}")

        # Stage 3: Parallel AI triage. No checkpoint/cache -- every row in
        # deduped_report is submitted and paid for exactly once per run. A prior
        # version retried automatically on unresolved rows because the checkpoint
        # made a retry free (only unfinished rows were re-fetched); without a
        # checkpoint a retry would resubmit and re-pay for every row, so this
        # stage now runs once and just reports unresolved rows as a warning --
        # rerun parallel_triage_agent.py yourself if you want to fill them in.
        before_triage_files = set(os.listdir(OUTPUT_TRIAGE_DIR)) if os.path.isdir(OUTPUT_TRIAGE_DIR) else set()
        stage3_elapsed, _ = run_stage(
            f"3/{total_stages} Parallel AI Triage (parallel_triage_agent.py)",
            [sys.executable, TRIAGE_SCRIPT, "--input", deduped_report, "--processor", args.processor],
        )
        after_triage_files = set(os.listdir(OUTPUT_TRIAGE_DIR)) if os.path.isdir(OUTPUT_TRIAGE_DIR) else set()
        new_files = sorted(after_triage_files - before_triage_files)
        if new_files:
            triaged_report = os.path.join(OUTPUT_TRIAGE_DIR, new_files[-1])
        else:
            triaged_report = newest_xlsx(OUTPUT_TRIAGE_DIR)

        counts = read_triage_insights_counts(triaged_report)
        unresolved = counts.get("ERROR", 0) + counts.get("NOT PROCESSED", 0)

        print(f"\n[Stage 3] Triaged report: {triaged_report}")
        if unresolved:
            print(
                f"  WARNING: {unresolved} row(s) are ERROR/NOT PROCESSED in this report. "
                "This report is not fully triaged."
            )
            print(f"  To retry (re-submits and re-pays for ALL rows, no checkpoint): "
                  f"python parallel_triage_agent.py --input \"{deduped_report}\" --processor {args.processor}")

        # Stage 4: mail drafting (pre-filter -> provenance -> drafting). Not fatal:
        # zero mail-worthy items this run is a normal outcome (see module docstring
        # in draft_morning_update.py -- most candidates are expected to fail the
        # on-topic/credibility/freshness gate), so a non-zero exit here must not
        # discard the already-good Stage 1-3 output or mark the run as failed.
        mail_draft = None
        mail_ok = None
        stage4_elapsed = 0.0
        if args.skip_mail_draft:
            print("\n[Stage 4] Skipped (--skip-mail-draft).")
        else:
            before_mail_files = set(os.listdir(OUTPUT_MAIL_DRAFT_DIR)) if os.path.isdir(OUTPUT_MAIL_DRAFT_DIR) else set()
            mail_cmd = [
                sys.executable, MAIL_SCRIPT,
                "--input", triaged_report,
                "--processor", args.processor,
            ]
            if proceed_with_draft:
                mail_cmd += ["--output-format", draft_output_format]
            else:
                mail_cmd += ["--dry-run"]
            if window_start:
                mail_cmd += ["--window-start", window_start]
            stage4_elapsed, mail_ok = run_stage(
                f"4/{total_stages} Mail Drafting (draft_morning_update.py)",
                mail_cmd,
                fatal=False,
            )
            after_mail_files = set(os.listdir(OUTPUT_MAIL_DRAFT_DIR)) if os.path.isdir(OUTPUT_MAIL_DRAFT_DIR) else set()
            new_mail_files = sorted(after_mail_files - before_mail_files)
            new_docx = [f for f in new_mail_files if f.endswith(".docx")]
            if new_docx:
                mail_draft = os.path.join(OUTPUT_MAIL_DRAFT_DIR, new_docx[-1])
            print(f"\n[Stage 4] Mail draft: {mail_draft or '(none produced -- see stage output above, e.g. zero mail-worthy items)'}")

        total_elapsed = time.time() - overall_start

        print("\n" + "=" * 70)
        if unresolved:
            print("  END-TO-END RUN FINISHED WITH UNRESOLVED ROWS")
        elif mail_ok is False:
            print("  END-TO-END RUN COMPLETE (NO MAIL DRAFT PRODUCED)")
        else:
            print("  END-TO-END RUN COMPLETE")
        print("=" * 70)
        print(f"  1. Extraction : {combined_report}  ({stage1_elapsed:.0f}s)")
        print(f"  2. Dedup      : {deduped_report}  ({stage2_elapsed:.0f}s)")
        print(f"  3. Triage     : {triaged_report}  ({stage3_elapsed:.0f}s)")
        if unresolved:
            print(f"                  WARNING: {unresolved} row(s) still ERROR/NOT PROCESSED -- see message above")
        if not args.skip_mail_draft:
            print(f"  4. Mail draft : {mail_draft or '(none produced)'}  ({stage4_elapsed:.0f}s)")
        print(f"  Total time    : {total_elapsed:.0f}s")
        print(f"  Log file      : {log_path}")
        print("=" * 70)

    except SystemExit as se:
        if isinstance(se.code, int):
            exit_code = se.code
        elif se.code is not None:
            print(f"\n{se.code}", file=sys.stderr)
            exit_code = 1
        print(f"\n[Logging] Run stopped early. Log saved to: {log_path}")
    except KeyboardInterrupt:
        exit_code = 130
        print(f"\n[Logging] Run interrupted by user. Log saved to: {log_path}")
    except Exception as exc:
        exit_code = 1
        import traceback
        traceback.print_exc()
        print(f"\n[Logging] Unhandled exception occurred: {exc}. Log saved to: {log_path}", file=sys.stderr)
    finally:
        if sys.stdin.isatty():
            try:
                _original_input("\nPress Enter to close this window...")
            except EOFError:
                pass
        try:
            log_file.flush()
            log_file.close()
        except Exception:
            pass
        sys.stdout = orig_stdout
        sys.stderr = orig_stderr
        builtins.input = _original_input
        if exit_code != 0:
            sys.exit(exit_code)


if __name__ == "__main__":
    main()
